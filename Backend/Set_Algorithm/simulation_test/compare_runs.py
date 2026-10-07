"""
compare_runs.py

Reads every simulate.py workbook in a folder and writes one summary
workbook (comparison_summary.xlsx) with a row per run, so all the
combinations can be compared side by side.

    python3 compare_runs.py comparison_results
"""

import os
import re
import sys
from collections import Counter

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
NAME_RATING = re.compile(r"(.*) \((\d+)\)$")


def summarise(path):
    wb = load_workbook(path, read_only=True)
    settings = {r[0]: r[1] for r in wb["Run Settings"].iter_rows(min_row=2, values_only=True) if r[0]}
    if "Priority" not in settings:
        raise ValueError("made by an old simulate.py (no 'Priority' row) -- re-run install_prototype4.sh")

    team_gaps, partner_gaps = Counter(), Counter()
    opp = opp_rep = part = part_rep = 0
    seen_p = seen_o = None
    day_now = None
    for row in wb["Round By Round Detail"].iter_rows(min_row=2, values_only=True):
        if row[0] != day_now:
            day_now, seen_p, seen_o = row[0], set(), set()
        (a1, r1), (a2, r2), (b1, r3), (b2, r4) = (NAME_RATING.match(c).groups() for c in row[3:7])
        r1, r2, r3, r4 = int(r1), int(r2), int(r3), int(r4)
        team_gaps[abs(r1 + r2 - r3 - r4)] += 1
        partner_gaps[max(abs(r1 - r2), abs(r3 - r4))] += 1
        for pair in ((a1, a2), (b1, b2)):
            key = frozenset(pair)
            part += 1
            part_rep += key in seen_p
            seen_p.add(key)
        for x in (a1, a2):
            for y in (b1, b2):
                key = frozenset((x, y))
                opp += 1
                opp_rep += key in seen_o
                seen_o.add(key)
    courts = sum(team_gaps.values())

    tally = list(wb["Final SitOut Tally"].iter_rows(min_row=2, values_only=True))
    ratios = [r[8] for r in tally if isinstance(r[8], (int, float))]
    most_in_day = max((r[7] for r in tally if isinstance(r[7], int)), default=0)

    pct = lambda n: round(100 * n / courts, 1) if courts else 0
    return [
        os.path.splitext(os.path.basename(path))[0],
        settings.get("Players attending per day"),
        settings.get("Courts"),
        settings.get("Rounds per session"),
        settings.get("Days simulated"),
        settings.get("Gender policy"),
        settings.get("Variety mode"),
        settings.get("Who-played-with-who memory"),
        settings.get("Max team rating gap"),
        settings.get("Max partner gap"),
        courts,
        round(100 * opp_rep / opp, 1) if opp else 0,
        round(100 * part_rep / part, 1) if part else 0,
        pct(sum(v for k, v in team_gaps.items() if k <= 1)),
        pct(sum(v for k, v in team_gaps.items() if k >= 3)),
        max(team_gaps) if team_gaps else 0,
        pct(sum(v for k, v in partner_gaps.items() if k >= 3)),
        max(partner_gaps) if partner_gaps else 0,
        (settings.get("Courts over team-gap limit") or 0) + (settings.get("Courts over partner-gap limit") or 0)
        + (settings.get("Courts over court-range limit") or 0),
        round(min(ratios), 4) if ratios else "no sit-outs",
        round(max(ratios), 4) if ratios else "no sit-outs",
        most_in_day,
        settings.get("Back-to-back sit-outs"),
    ]


HEADERS = [
    "Run", "Players/day", "Courts", "Rounds", "Days", "Gender", "Variety mode", "Memory",
    "Max team gap", "Max partner gap", "Courts played",
    "Repeat opponents % (same session)", "Repeat partners % (same session)",
    "Team gap 0-1 %", "Team gap 3+ %", "Worst team gap",
    "Partners 3+ apart %", "Worst partner gap", "Courts over a limit",
    "Sit-out ratio min", "Sit-out ratio max", "Most sit-outs in a day", "Back-to-back sit-outs",
]


def main():
    folder = sys.argv[1] if len(sys.argv) > 1 else "comparison_results"
    files = sorted(f for f in os.listdir(folder)
                   if f.endswith(".xlsx") and not f.startswith(("comparison_summary", "~$")))
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"
    ws.append(HEADERS)
    skipped = []
    for f in files:
        try:
            ws.append(summarise(os.path.join(folder, f)))
        except Exception as e:  # keep going; report at the end
            skipped.append(f"{f}: {e}")
    for cell in ws[1]:
        cell.font = BOLD
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = FONT
    ws.freeze_panes = "B2"
    for col in ws.columns:
        width = max(len(str(c.value)) for c in col if c.value is not None)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 44)
    out = os.path.join(folder, "comparison_summary.xlsx")
    wb.save(out)
    print(f"Wrote {out} ({len(files) - len(skipped)} runs)")
    for s in skipped:
        print("  skipped", s)


if __name__ == "__main__":
    main()

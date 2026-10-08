"""
simulate.py  (runs against Prototype4)

Minimal-input, long-run simulation against your REAL roster.csv and the
REAL algorithm in the sibling ../Prototype4 folder (not a reimplementation).

Run it with just:
    python3 simulate.py
and answer the prompts (every one has a default -- press Enter to accept).

Produces ONE workbook, simulation_results_prototype4.xlsx (the Prototype3
version writes simulation_results.xlsx, so the two never overwrite each other):
  Run Settings           -- what this run used (algorithm folder, gender
                            policy, variety mode, skill limits, memory...)
  Summary Checkpoints    -- every player's cumulative sit-out count and
                            ratio at the end of every simulated day.
  Final SitOut Tally     -- one row per player, final totals, ranked, plus the
                            most times they sat out in a single day.
  Skill Balance Summary  -- one row per day: average/worst team rating gap,
                            worst court range, courts over the limits.
  Round By Round Detail  -- every court, every round, every day.

The previous Prototype3 version of this script is kept next to it as
simulate_prototype3.py (if you ran the installer), so you can run both on
the same settings and compare the two workbooks.
"""

import csv
import os
import random
import sys
from collections import deque

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
PROTOTYPE_DIR = os.path.join(HERE, "..", "Prototype4")
ROSTER_PATH = os.path.join(HERE, "roster.csv")
OUTPUT_PATH = os.path.join(HERE, "simulation_results_prototype4.xlsx")
SEED = 2026

# Skill settings come from the variety mode you pick (algorithm.MODE_PROFILES).
# Set any of these to a number to override that one setting for every mode.
MAX_TEAM_DIFF = None
MAX_COURT_SPREAD = None
MAX_PARTNER_GAP = None
IDEAL_TEAM_DIFF = None

FONT = Font(name="Arial", size=10)
HEADER_FONT = Font(name="Arial", size=10, bold=True)
OVER_LIMIT_FILL = PatternFill("solid", start_color="FCE4D6")


def _prompt_int(label, default):
    raw = input(f"{label} [{default}]: ").strip()
    if raw == "":
        return default
    try:
        return int(raw)
    except ValueError:
        print(f"  Not a whole number, using default ({default}).")
        return default


def _prompt_choice(label, choices, default):
    desc = ", ".join(f"{k}={v}" for k, v in choices.items())
    raw = input(f"{label} ({desc}) [{default}]: ").strip()
    if raw == "":
        return default
    try:
        value = int(raw)
    except ValueError:
        print(f"  Not a valid choice, using default ({default}).")
        return default
    if value not in choices:
        print(f"  Not a valid choice, using default ({default}).")
        return default
    return value


def load_roster(path):
    rows = []
    with open(path, newline="") as f:
        for row in csv.DictReader(f):
            rows.append({
                "player_id": row["player_id"].strip(),
                "name": row["name"].strip(),
                "gender": row["gender"].strip().upper(),
                "rating": int(row["rating"]),
                "starting_sit_out_count": int(row["starting_sit_out_count"]) if row.get("starting_sit_out_count", "").strip() else 0,
            })
    return rows


def main():
    sys.path.insert(0, os.path.abspath(PROTOTYPE_DIR))
    try:
        import algorithm
        from helpers import Player
    except ImportError as e:
        print(f"Couldn't import the algorithm from {os.path.abspath(PROTOTYPE_DIR)} ({e}).\n"
              f"simulate.py expects a Prototype4/ folder next to simulation_test/.")
        sys.exit(1)
    if not hasattr(algorithm, "RATIO_WINDOW_ROUNDS"):
        print(f"{os.path.abspath(PROTOTYPE_DIR)} isn't the latest Prototype4. "
              f"Re-run install_prototype4.sh.")
        sys.exit(1)

    window = algorithm.RATIO_WINDOW_ROUNDS

    print("=== TennisSetMaker simulation (Prototype4) ===")
    print(f"Roster: {ROSTER_PATH}")
    print(f"Sit-outs: everyone sits once before anyone sits twice, no back-to-back, "
          f"ties by ratio over each player's last {window} rounds\n")

    rounds_per_session = _prompt_int("Rounds per session (a normal club day)", default=6)
    gender_labels = {0: "open/any gender", 1: "enforce mixed pairs", 2: "same-gender courts"}
    mixed_mode = _prompt_choice("Gender policy", gender_labels, default=0)
    variety_labels = {1: "hard avoid repeats (rotation first, gaps up to 4)",
                      2: "soft (skill first, gaps up to 2)",
                      3: "round-robin (middle: gaps up to 2, avoid anyone met today)"}
    variety_mode = _prompt_choice("Variety mode", variety_labels, default=2)
    overrides = {k: v for k, v in dict(max_team_diff=MAX_TEAM_DIFF, max_court_spread=MAX_COURT_SPREAD,
                                        max_partner_gap=MAX_PARTNER_GAP, ideal_team_diff=IDEAL_TEAM_DIFF).items()
                 if v is not None}
    profile = algorithm.skill_profile(variety_mode)
    profile.update(overrides)
    max_team_diff, max_partner_gap = profile["max_team_diff"], profile["max_partner_gap"]
    max_court_spread, ideal_team_diff = profile["max_court_spread"], profile["ideal_team_diff"]
    print(f"  -> {'rotation' if profile['rotation_first'] else 'skill'} first; team gap <= {max_team_diff}, "
          f"partners <= {max_partner_gap} apart"
          + (f", court range <= {max_court_spread}" if max_court_spread is not None else "")
          + (f", team gap {ideal_team_diff} preferred" if ideal_team_diff is not None else "") + "\n")
    memory_labels = {1: "reset every session", 2: "carry across all sessions"}
    memory_mode = _prompt_choice("Who-played-with-who memory", memory_labels, default=1)

    roster_rows = load_roster(ROSTER_PATH)
    full_pool_size = len(roster_rows)

    # Round generation scores every possible foursome from the remaining
    # pool (C(n,4)), so it slows down sharply past ~40 players per round.
    num_courts = _prompt_int("Number of courts your club has available", default=6)
    if num_courts > 12:
        print(f"  Note: {num_courts * 4} players/round -- this could take minutes per day.")
    num_to_play = num_courts * 4

    players_per_day = _prompt_int(
        f"Players attending each day (drawn at random from your {full_pool_size}-player roster)",
        default=min(full_pool_size, num_to_play + 8),
    )
    if players_per_day > full_pool_size:
        print(f"  Only {full_pool_size} players in the roster -- capping to that.")
        players_per_day = full_pool_size
    if players_per_day <= num_to_play:
        print(f"  Note: {players_per_day} attending <= {num_to_play} court spots -- nobody will sit out.")
    num_days = _prompt_int("Number of days/sessions to simulate", default=100)

    print(
        f"\n{players_per_day} of {full_pool_size} players drawn per day -> {num_courts} court(s), "
        f"{num_to_play} play per round.\n"
        f"Simulating {num_days} days x {rounds_per_session} rounds = {num_days * rounds_per_session} rounds...\n"
    )

    attendance_rng = random.Random(SEED)
    rng = random.Random(SEED + 1)
    all_players = {
        r["player_id"]: Player(
            name=r["name"], gender=r["gender"], rating=r["rating"],
            player_id=r["player_id"], sit_out_count=r["starting_sit_out_count"],
        )
        for r in roster_rows
    }
    # Kept separate from Player.sessions_played, which the algorithm itself
    # increments once per round played.
    days_attended = {pid: 0 for pid in all_players}
    # Each player's last `window` rounds attended (1 = sat out, 0 = played),
    # standing in for what the backend would pass as recent_sit_outs/recent_rounds.
    recent = {pid: deque(maxlen=window) for pid in all_players}
    most_in_a_day = {pid: 0 for pid in all_players}
    back_to_back = 0
    checkpoint_rows, skill_rows, round_rows = [], [], []
    global_round = 0
    total_courts = courts_over_diff = courts_over_spread = courts_over_partner = 0

    for day in range(1, num_days + 1):
        if memory_mode == 1:
            for p in all_players.values():
                p.partners_played.clear()
                p.opponents_played.clear()
        attendees = attendance_rng.sample(list(all_players.values()), players_per_day)
        for p in attendees:
            days_attended[p.player_id] += 1
            p.session_sit_outs = 0       # new session
            p.consecutive_sit_outs = 0
        day_diffs, day_spreads = [], []
        day_over = 0
        for round_in_day in range(1, rounds_per_session + 1):
            global_round += 1
            for p in attendees:
                p.recent_sit_outs = sum(recent[p.player_id])
                p.recent_rounds = len(recent[p.player_id])
            playing, sitting_out = algorithm.select_round_players(attendees, num_to_play, rng)
            court_results = algorithm.assign_courts_for_round(
                playing, num_courts, mixed_mode, variety_mode, **overrides,
            )

            sitting_ids = {id(p) for p in sitting_out}
            for p in attendees:
                if id(p) in sitting_ids:
                    back_to_back += p.consecutive_sit_outs > 0
                    p.sit_out_count += 1
                    p.session_sit_outs += 1
                    p.last_sit_out_round = global_round
                    p.consecutive_sit_outs += 1
                    recent[p.player_id].append(1)
                    most_in_a_day[p.player_id] = max(most_in_a_day[p.player_id], p.session_sit_outs)
                else:
                    p.consecutive_sit_outs = 0
                    recent[p.player_id].append(0)

            for r in court_results:
                ta, tb = r["team_a"], r["team_b"]
                diff = abs(r["team_a_sum"] - r["team_b_sum"])
                spread = algorithm.raw_skill_spread(tuple(r["players"]))
                pgap = algorithm.raw_partner_gap(ta, tb)
                over_diff = max_team_diff is not None and diff > max_team_diff
                over_spread = max_court_spread is not None and spread > max_court_spread
                over_partner = max_partner_gap is not None and pgap > max_partner_gap
                over = over_diff or over_spread or over_partner
                total_courts += 1
                courts_over_diff += over_diff
                courts_over_spread += over_spread
                courts_over_partner += over_partner
                day_over += over
                day_diffs.append(diff)
                day_spreads.append(spread)
                round_rows.append([
                    day, round_in_day, r["court"],
                    f"{ta[0].name} ({ta[0].rating})", f"{ta[1].name} ({ta[1].rating})",
                    f"{tb[0].name} ({tb[0].rating})", f"{tb[1].name} ({tb[1].rating})",
                    r["team_a_sum"], r["team_b_sum"], diff, pgap, spread,
                    "".join(p.gender[0] for p in ta) + " v " + "".join(p.gender[0] for p in tb),
                    "YES" if over else "",
                ])

        for p in all_players.values():
            eligible = days_attended[p.player_id] * rounds_per_session
            ratio = (p.sit_out_count / eligible) if eligible else None
            checkpoint_rows.append([
                day, p.name, p.gender, p.rating, p.sit_out_count, eligible,
                round(ratio, 4) if ratio is not None else "no_data_yet",
            ])

        if day_diffs:
            skill_rows.append([day, round(sum(day_diffs) / len(day_diffs), 3), max(day_diffs),
                               max(day_spreads), day_over, len(day_diffs)])

        if day % 10 == 0:
            print(f"  ...day {day}/{num_days} done", flush=True)

    final_rows = []
    for p in all_players.values():
        attended = days_attended[p.player_id]
        eligible = attended * rounds_per_session
        ratio = (p.sit_out_count / eligible) if eligible else None
        final_rows.append([p.name, p.gender, p.rating, p.sit_out_count, attended, eligible,
                           most_in_a_day[p.player_id], ratio])
    final_rows.sort(key=lambda row: (row[-1] is None, -(row[-1] or 0)))

    settings = [
        ["Algorithm folder", "Prototype4"],
        ["Gender policy", f"{mixed_mode} = {gender_labels[mixed_mode]}"],
        ["Variety mode", f"{variety_mode} = {variety_labels[variety_mode]}"],
        ["Who-played-with-who memory", memory_labels[memory_mode]],
        ["Priority", "rotation first" if profile["rotation_first"] else "skill first"],
        ["Max team rating gap", max_team_diff],
        ["Ideal team rating gap", ideal_team_diff if ideal_team_diff is not None else "-"],
        ["Max partner gap", max_partner_gap],
        ["Max court range (highest - lowest)", max_court_spread if max_court_spread is not None else "-"],
        ["Sit-out rules", "everyone sits once before anyone sits twice; no back-to-back; "
                          f"ties by ratio over last {window} rounds"],
        ["Courts", num_courts],
        ["Players attending per day", f"{players_per_day} of {full_pool_size}"],
        ["Rounds per session", rounds_per_session],
        ["Days simulated", num_days],
        ["Seed", SEED],
        ["", ""],
        ["Courts played", total_courts],
        ["Courts over team-gap limit", courts_over_diff],
        ["Courts over partner-gap limit", courts_over_partner],
        ["Courts over court-range limit", courts_over_spread],
        ["Back-to-back sit-outs", back_to_back],
    ]
    _write_workbook(settings, checkpoint_rows, final_rows, skill_rows, round_rows)

    ratios = [row[-1] for row in final_rows if row[-1] is not None]
    if ratios:
        print(f"Done. Sit-out ratio range: min={min(ratios):.3f}, max={max(ratios):.3f}, mean={sum(ratios)/len(ratios):.3f}")
    print(f"Courts over a limit -- team gap: {courts_over_diff}, partner gap: {courts_over_partner}, "
          f"court range: {courts_over_spread} (of {total_courts}); back-to-back sit-outs: {back_to_back}")
    print(f"Wrote: {OUTPUT_PATH}")


def _style_header(ws, num_cols):
    for col in range(1, num_cols + 1):
        ws.cell(row=1, column=col).font = HEADER_FONT
    ws.freeze_panes = "A2"


def _write_workbook(settings, checkpoint_rows, final_rows, skill_rows, round_rows):
    wb = Workbook()

    ws0 = wb.active
    ws0.title = "Run Settings"
    ws0.append(["Setting", "Value"])
    for row in settings:
        ws0.append(row)
    _style_header(ws0, 2)

    ws1 = wb.create_sheet("Summary Checkpoints")
    ws1.append(["Day", "Player", "Gender", "Rating", "Cumulative_SitOut_Count", "Rounds_Eligible_So_Far", "SitOut_Ratio"])
    for row in checkpoint_rows:
        ws1.append(row)
    _style_header(ws1, 7)

    ws2 = wb.create_sheet("Final SitOut Tally")
    ws2.append(["Rank", "Player", "Gender", "Rating", "Total_SitOut_Count", "Days_Attended", "Total_Rounds_Eligible",
                "Most_SitOuts_In_A_Day", "SitOut_Ratio"])
    for rank, row in enumerate(final_rows, start=1):
        name, gender, rating, sit_out_count, attended, eligible, most, ratio = row
        ws2.append([rank, name, gender, rating, sit_out_count, attended, eligible, most,
                    round(ratio, 4) if ratio is not None else "never_drawn"])
    _style_header(ws2, 9)
    last_row = len(final_rows) + 1
    ws2["K1"], ws2["L1"] = "Stat", "SitOut_Ratio"
    ws2["K1"].font = HEADER_FONT
    ws2["L1"].font = HEADER_FONT
    for i, (label, fn) in enumerate([("Min", "MIN"), ("Max", "MAX"), ("Average", "AVERAGE")], start=2):
        ws2.cell(row=i, column=11, value=label)
        ws2.cell(row=i, column=12, value=f"={fn}(I2:I{last_row})")

    ws3 = wb.create_sheet("Skill Balance Summary")
    ws3.append(["Day", "Avg_Team_Gap", "Worst_Team_Gap", "Worst_Court_Range", "Courts_Over_Limit", "Courts_Played"])
    for row in skill_rows:
        ws3.append(row)
    _style_header(ws3, 6)

    ws4 = wb.create_sheet("Round By Round Detail")
    ws4.append(["Day", "Round", "Court", "TeamA_Player1", "TeamA_Player2", "TeamB_Player1", "TeamB_Player2",
                "TeamA_Sum", "TeamB_Sum", "Team_Gap", "Partner_Gap", "Court_Range", "Genders", "Over_Limit"])
    for row in round_rows:
        ws4.append(row)
        if row[-1] == "YES":
            for cell in ws4[ws4.max_row]:
                cell.fill = OVER_LIMIT_FILL
    _style_header(ws4, 14)

    for ws in (ws0, ws1, ws2, ws3, ws4):
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = FONT
        for col_cells in ws.columns:
            width = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max(width + 2, 10), 36)

    wb.save(OUTPUT_PATH)


if __name__ == "__main__":
    main()
